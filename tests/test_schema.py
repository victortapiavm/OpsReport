from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.anomalies import detect_anomalies
from src.exports import build_excel_export
from src.metrics import calculate_kpis
from src.narrative import generate_management_narrative
from src.profiling import profile_dataframe
from src.schema import (
    SEMANTIC_FIELDS,
    infer_schema,
    resolve_column,
    suggested_mapping,
    validate_column_mapping,
)
from src.validation import assess_data_quality


def test_infers_common_english_headers_with_high_confidence() -> None:
    df = pd.DataFrame(
        {
            "Order Number": ["A", "B"],
            "Transaction Date": ["2026-09-01", "2026-09-02"],
            "Net Sales": [100, 120],
            "COGS": [60, 70],
            "Units Sold": [1, 2],
            "Order Status": ["Completed", "Cancelled"],
            "Market Region": ["North", "South"],
            "Sales Channel": ["Online", "Store"],
            "Product Family": ["Tools", "Parts"],
            "Processing Hours": [2.0, 3.5],
        }
    )

    inferences = infer_schema(df)
    mapping = suggested_mapping(inferences)

    assert set(mapping) == set(SEMANTIC_FIELDS)
    assert mapping["order_id"] == "Order Number"
    assert mapping["revenue"] == "Net Sales"
    assert mapping["channel"] == "Sales Channel"
    assert mapping["processing_time_hours"] == "Processing Hours"
    assert all(item.confidence == "high" for item in inferences.values())


def test_infers_spanish_headers_accent_and_separator_insensitive() -> None:
    df = pd.DataFrame(
        {
            "N° Pedido": ["A", "B"],
            "Fecha Operación": ["2026-09-01", "2026-09-02"],
            "Facturación": [100, 120],
            "Costo Total": [60, 70],
            "Unidades": [1, 2],
            "Estado Pedido": ["Completado", "Cancelado"],
            "Zona": ["Centro", "Sur"],
            "Canal Venta": ["Web", "Tienda"],
            "Categoría Producto": ["Herramientas", "Repuestos"],
            "Horas Proceso": [2.0, 3.5],
        }
    )

    mapping = suggested_mapping(infer_schema(df))

    assert mapping == {
        "order_id": "N° Pedido",
        "date": "Fecha Operación",
        "revenue": "Facturación",
        "cost": "Costo Total",
        "quantity": "Unidades",
        "status": "Estado Pedido",
        "region": "Zona",
        "channel": "Canal Venta",
        "category": "Categoría Producto",
        "processing_time_hours": "Horas Proceso",
    }


def test_ambiguous_revenue_headers_require_confirmation() -> None:
    df = pd.DataFrame({"Sales": [10, 20], "Revenue": [11, 21], "notes": ["a", "b"]})

    inferences = infer_schema(df)

    assert inferences["revenue"].column in {"Sales", "Revenue"}
    assert inferences["revenue"].confidence == "low"
    assert inferences["revenue"].requires_confirmation is True
    assert set(inferences["revenue"].alternatives) == (
        {"Sales", "Revenue"} - {inferences["revenue"].column}
    )
    assert "revenue" not in suggested_mapping(inferences)


def test_missing_fields_stay_unmapped_and_explicit_none_disables_fallback() -> None:
    df = pd.DataFrame(
        {
            "revenue": [10, 20],
            "created_at": ["2026-09-01 08:00", "2026-09-02 08:00"],
            "completed_at": ["2026-09-01 10:00", "2026-09-02 11:00"],
            "notes": ["a", "b"],
        }
    )

    inferences = infer_schema(df)
    automatic = suggested_mapping(inferences)
    explicit = {field: automatic.get(field) for field in SEMANTIC_FIELDS}

    assert automatic == {"revenue": "revenue"}
    assert explicit["order_id"] is None
    assert resolve_column(df, "revenue", explicit) == "revenue"

    explicit["revenue"] = None
    assert resolve_column(df, "revenue", explicit) is None
    metrics = calculate_kpis(df, column_map=explicit)
    assert metrics["revenue"] is None
    assert metrics["average_processing_time_hours"] is None


def test_user_override_wins_and_duplicate_source_mapping_is_rejected() -> None:
    df = pd.DataFrame(
        {
            "Gross Sales": [100, 200],
            "Net Billing": [80, 160],
            "Order Ref": ["A", "B"],
        }
    )
    mapping = {"revenue": "Net Billing", "order_id": "Order Ref"}

    assert resolve_column(df, "revenue", mapping) == "Net Billing"
    assert calculate_kpis(df, column_map=mapping)["revenue"] == 240.0
    assert validate_column_mapping(df, mapping) == []

    errors = validate_column_mapping(
        df,
        {"revenue": "Net Billing", "cost": "Net Billing"},
    )
    assert len(errors) == 1
    assert "Cada columna debe tener un solo rol" in errors[0]


def test_normalized_header_collision_is_not_silently_resolved() -> None:
    df = pd.DataFrame({"Sales": [10, 20], "sales": [11, 21]})

    assert resolve_column(df, "revenue") is None


def test_manual_mapping_drives_full_analysis_and_type_diagnostics() -> None:
    df = pd.DataFrame(
        {
            "Ref": ["A", "B", "C", "D", "E"],
            "When": ["2026-09-01", "2026-09-02", "bad-date", "2026-09-04", "2026-09-05"],
            "Billings": [100, 102, 98, 101, 1000],
            "Spend": [60, 61, 58, 60, 400],
            "Pieces": [1, 2, "bad-number", 1, 3],
            "Lifecycle": ["Done", "Done", "Cancelled", "Done", "Done"],
            "Market": ["N", "N", "S", "S", "S"],
            "Line": ["A", "A", "B", "B", "B"],
            "Cycle": [2, 2, 2, 2, 20],
        }
    )
    mapping = {
        "order_id": "Ref",
        "date": "When",
        "revenue": "Billings",
        "cost": "Spend",
        "quantity": "Pieces",
        "status": "Lifecycle",
        "region": "Market",
        "category": "Line",
        "processing_time_hours": "Cycle",
    }

    quality = assess_data_quality(df, column_map=mapping)
    metrics = calculate_kpis(df, column_map=mapping)
    anomalies = detect_anomalies(df, column_map=mapping)
    profile = profile_dataframe(df)
    narrative = generate_management_narrative(
        df,
        profile=profile,
        metrics=metrics,
        anomalies=anomalies,
        column_map=mapping,
    )

    assert quality["stats"]["invalid_numeric_values"] == 1
    assert quality["stats"]["invalid_date_values"] == 1
    assert metrics["order_count"] == 5
    assert metrics["revenue"] == 1401.0
    assert metrics["units"] == 7.0
    assert metrics["cancellation_rate"] == pytest.approx(20.0)
    assert any(item["canonical_column"] == "revenue" for item in anomalies)
    assert any(item["canonical_column"] == "processing_time_hours" for item in anomalies)
    assert "01-09-2026 a 05-09-2026" in narrative
    assert "$1.401" in narrative

    workbook_bytes = build_excel_export(
        df,
        profile=profile,
        metrics=metrics,
        anomalies=anomalies,
        context={"column_mapping": mapping},
        column_map=mapping,
    )
    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True)
    assert {"Resumen", "Datos", "Calidad", "Métricas", "Anomalías", "Contexto"}.issubset(
        workbook.sheetnames
    )
