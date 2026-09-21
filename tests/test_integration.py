from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from src.anomalies import detect_anomalies
from src.exports import build_excel_export
from src.ingestion import load_data
from src.metrics import calculate_kpis
from src.narrative import generate_management_narrative
from src.profiling import profile_dataframe
from src.validation import assess_data_quality


SAMPLE_PATH = Path(__file__).parents[1] / "data" / "sample_operations.csv"


def test_bundled_sample_runs_end_to_end() -> None:
    dataframe = load_data(SAMPLE_PATH, filename=SAMPLE_PATH.name)

    assert dataframe.shape == (1500, 10)
    assert {
        "order_id",
        "date",
        "region",
        "channel",
        "category",
        "status",
        "units",
        "revenue",
        "cost",
        "processing_hours",
    }.issubset(dataframe.columns)

    profile = profile_dataframe(dataframe)
    quality = assess_data_quality(dataframe)
    metrics = calculate_kpis(dataframe)
    anomalies = detect_anomalies(dataframe)
    narrative = generate_management_narrative(
        dataframe,
        profile=profile,
        metrics=metrics,
        anomalies=anomalies,
    )

    assert profile["missing_cells"] > 0
    assert quality["duplicate_count"] == 10
    assert 0 < quality["score"] < 100
    assert metrics["order_count"] == 1490
    assert metrics["revenue"] is not None and metrics["revenue"] > 0
    assert metrics["gross_margin"] is not None and metrics["gross_margin"] > 0
    assert metrics["average_processing_time_hours"] is not None
    assert anomalies
    assert "Se analizaron 1.500 registros" in narrative

    workbook_bytes = build_excel_export(
        dataframe,
        profile=profile,
        metrics=metrics,
        anomalies=anomalies,
        context={"source_file": SAMPLE_PATH.name},
    )
    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True)
    assert {"Resumen", "Datos", "Calidad", "Métricas", "Anomalías", "Contexto"}.issubset(
        workbook.sheetnames
    )
