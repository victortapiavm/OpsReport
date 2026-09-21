"""Excel export helpers for a portable OpsReport analysis package."""

from __future__ import annotations

from io import BytesIO
from typing import Any, Mapping

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.narrative import generate_management_narrative
from src.profiling import categorical_summary, missingness_table, numeric_summary, profile_dataframe

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def build_excel_export(
    dataframe: pd.DataFrame,
    profile: Mapping[str, Any] | None = None,
    metrics: Mapping[str, Any] | pd.DataFrame | None = None,
    anomalies: pd.DataFrame | list[Mapping[str, Any]] | None = None,
    context: Mapping[str, Any] | None = None,
    column_map: Mapping[str, str | None] | None = None,
) -> bytes:
    """Build a styled XLSX with raw data, context and analysis tables."""
    resolved_profile = dict(profile or profile_dataframe(dataframe))
    output = BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        _summary_table(dataframe, resolved_profile, metrics, anomalies, column_map).to_excel(
            writer, sheet_name="Resumen", index=False
        )
        dataframe.to_excel(writer, sheet_name="Datos", index=False)
        missingness_table(dataframe).to_excel(writer, sheet_name="Calidad", index=False)
        numeric_summary(dataframe).to_excel(writer, sheet_name="Numéricos", index=False)
        categorical_summary(dataframe).to_excel(writer, sheet_name="Categóricos", index=False)

        context_frame = _mapping_table(context or {}, "campo", "valor")
        if context_frame.empty:
            context_frame = pd.DataFrame([{"campo": "nota", "valor": "Sin contexto adicional."}])
        context_frame.to_excel(writer, sheet_name="Contexto", index=False)

        metrics_frame = _coerce_table(metrics, "indicador", "valor")
        if not metrics_frame.empty:
            metrics_frame.to_excel(writer, sheet_name="Métricas", index=False)

        anomaly_frame = _coerce_anomalies(anomalies)
        if not anomaly_frame.empty:
            anomaly_frame.to_excel(writer, sheet_name="Anomalías", index=False)

        _style_workbook(writer.book)

    return output.getvalue()


export_excel = build_excel_export


def _summary_table(
    dataframe: pd.DataFrame,
    profile: Mapping[str, Any],
    metrics: Mapping[str, Any] | pd.DataFrame | None,
    anomalies: pd.DataFrame | list[Mapping[str, Any]] | None,
    column_map: Mapping[str, str | None] | None,
) -> pd.DataFrame:
    narrative = generate_management_narrative(
        dataframe,
        profile=profile,
        metrics=metrics,
        anomalies=anomalies,
        column_map=column_map,
    )
    rows: list[dict[str, Any]] = [
        {"sección": "Cobertura", "detalle": f"{profile.get('rows', len(dataframe))} registros × {profile.get('columns', len(dataframe.columns))} columnas"},
        {"sección": "Duplicados", "detalle": int(profile.get("duplicate_rows", 0))},
        {"sección": "Celdas vacías (%)", "detalle": float(profile.get("missing_pct", 0.0))},
    ]
    rows.extend({"sección": "Narrativa", "detalle": paragraph} for paragraph in narrative.split("\n\n"))
    return pd.DataFrame(rows)


def _mapping_table(mapping: Mapping[str, Any], key_name: str, value_name: str) -> pd.DataFrame:
    rows = [{key_name: str(key), value_name: _excel_safe(value)} for key, value in mapping.items()]
    return pd.DataFrame(rows, columns=[key_name, value_name])


def _coerce_table(
    value: Mapping[str, Any] | pd.DataFrame | None,
    key_name: str,
    value_name: str,
) -> pd.DataFrame:
    if value is None:
        return pd.DataFrame(columns=[key_name, value_name])
    if isinstance(value, pd.DataFrame):
        return value.copy()
    if isinstance(value, Mapping):
        return _mapping_table(value, key_name, value_name)
    return pd.DataFrame(columns=[key_name, value_name])


def _coerce_anomalies(anomalies: pd.DataFrame | list[Mapping[str, Any]] | None) -> pd.DataFrame:
    if anomalies is None:
        return pd.DataFrame()
    if isinstance(anomalies, pd.DataFrame):
        return anomalies.copy()
    return pd.DataFrame(anomalies)


def _excel_safe(value: Any) -> Any:
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    if isinstance(value, (list, tuple, set)):
        return ", ".join(map(str, value))
    if isinstance(value, Mapping):
        return "; ".join(f"{key}={item}" for key, item in value.items())
    return str(value)


def _style_workbook(workbook: Any) -> None:
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        if sheet.max_row >= 1 and sheet.max_column >= 1:
            for cell in sheet[1]:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(vertical="center")
            if sheet.max_row > 1:
                sheet.auto_filter.ref = sheet.dimensions

        for column_index, column_cells in enumerate(sheet.columns, start=1):
            values = [str(cell.value) for cell in column_cells if cell.value is not None]
            max_length = max((len(value) for value in values), default=8)
            sheet.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 10), 42)
        if sheet.title == "Resumen":
            sheet.column_dimensions["A"].width = 22
            sheet.column_dimensions["B"].width = 90
            for cell in sheet["B"]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
